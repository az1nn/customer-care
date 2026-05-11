from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
INPUT_FILE = ROOT / "slides.example.json"
OUTPUT_DIR = ROOT / "output"
OUTPUT_FILE = OUTPUT_DIR / "desenvolvimento-assistido-por-ia-az1nn.pptx"


def load_payload() -> dict:
    with INPUT_FILE.open("r", encoding="utf-8") as file:
        return json.load(file)


def build_presentation(payload: dict) -> Presentation:
    presentation = Presentation()

    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    cover_title = cover.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(12), Inches(1.3)).text_frame
    cover_title.text = payload.get("title", "Apresentação")
    cover_title.paragraphs[0].font.size = Pt(34)
    cover_title.paragraphs[0].font.bold = True

    cover_subtitle = cover.shapes.add_textbox(Inches(0.7), Inches(3.5), Inches(12), Inches(0.8)).text_frame
    cover_subtitle.text = payload.get("subtitle", "")
    cover_subtitle.paragraphs[0].font.size = Pt(20)

    for slide_data in payload.get("slides", []):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_data.get("title", "Sem título")
        body = slide.placeholders[1].text_frame
        body.clear()

        bullets = slide_data.get("bullets", [])
        for index, bullet in enumerate(bullets):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(20)

    return presentation


def main() -> None:
    payload = load_payload()
    presentation = build_presentation(payload)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT_FILE.as_posix())
    print(f"Arquivo gerado: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
